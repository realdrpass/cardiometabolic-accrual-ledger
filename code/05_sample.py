#!/usr/bin/env python3
"""Step 5: draw the 10 percent hand-check sample and build the fill-in form.

Plain-language note: this script picks about one trial in ten at random (with a
fixed seed so anyone can draw the same sample) and makes a spreadsheet for the
author. For each trial the form shows what the ledger says, a link to the
registry page, and blank yellow cells where the author types what the registry
page says. A formula marks each pair OK or CHECK. The Summary tab totals the
agreement rate. Open the spreadsheet, not this file.

Output: data/processed/handcheck_sample.csv
        report/handcheck_form.xlsx
"""
import math
import pathlib

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parents[1]
P = ROOT / "data" / "processed"
SEED = 20261013
FRACTION = 0.10

STOPPED = ("TERMINATED", "WITHDRAWN", "SUSPENDED")


def stratum(s):
    if s == "COMPLETED":
        return "completed"
    if s in STOPPED:
        return "stopped"
    return "ongoing_or_unknown"


# (ledger column, label shown to author, where to look on the registry page, code list or None)
STATUS_CODES = [
    ("COMPLETED", "Completed"), ("TERMINATED", "Terminated"), ("WITHDRAWN", "Withdrawn"),
    ("SUSPENDED", "Suspended"), ("RECRUITING", "Recruiting"),
    ("ACTIVE_NOT_RECRUITING", "Active, not recruiting"), ("NOT_YET_RECRUITING", "Not yet recruiting"),
    ("ENROLLING_BY_INVITATION", "Enrolling by invitation"), ("UNKNOWN", "Unknown status"),
]
PHASE_CODES = [
    ("PHASE2", "Phase 2"), ("PHASE3", "Phase 3"), ("PHASE4", "Phase 4"),
    ("PHASE1/PHASE2", "Phase 1 and Phase 2 both listed"), ("PHASE2/PHASE3", "Phase 2 and Phase 3 both listed"),
]
ENR_TYPE_CODES = [("ACTUAL", "(Actual) after the number"), ("ESTIMATED", "(Estimated) after the number")]
SPONSOR_CODES = [("INDUSTRY", "Lead sponsor is a company"),
                 ("other", "University, hospital, government agency, network or individual")]
YN_CODES = [("Y", "Yes"), ("N", "No")]
STOP_CODES = [
    ("poor_enrollment", "Poor enrollment"), ("business_sponsor", "Business or sponsor decision"),
    ("lack_of_benefit", "Lack of benefit"), ("safety", "Safety"),
    ("design_logistics", "Study design or logistics"), ("other_or_not_stated", "Other or not stated"),
    ("n/a", "Trial did not stop early (not Terminated, Withdrawn or Suspended)"),
]
CHECKS = [
    ("overall_status", "Status", "Top of page, next to the status label", STATUS_CODES),
    ("enrollment_reported", "Enrollment (number)", "Study Details > Study Design > Enrollment", None),
    ("enrollment_type", "Enrollment type", "Same line: (Actual) or (Estimated)", ENR_TYPE_CODES),
    ("phase", "Phase", "Study Details > Study Design > Phase", PHASE_CODES),
    ("sponsor_class", "Sponsor type", "'Sponsor' name; is it a company?", SPONSOR_CODES),
    ("us_site", "At least one US site?", "Contacts and Locations section", YN_CODES),
    ("in_scope_condition", "A trial of one of the 5 diseases?",
     "Conditions list and title; compare with CODING_RULES.md Part A", YN_CODES),
    ("stop_category", "Stop-reason group", "'Why Stopped' line; compare with CODING_RULES.md Part B", STOP_CODES),
]

YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREY = PatternFill("solid", fgColor="EDEDED")
HEAD = PatternFill("solid", fgColor="D9E1F2")
F = "Arial"


def main():
    L = pd.read_csv(P / "ledger.csv", keep_default_na=False)
    L["stratum"] = L.overall_status.map(stratum)
    n_target = math.ceil(len(L) * FRACTION)
    parts = []
    for name, grp in L.groupby("stratum"):
        k = round(len(grp) * FRACTION)
        parts.append(grp.sample(n=k, random_state=SEED))
    S = pd.concat(parts)
    if len(S) < n_target:  # top up rounding shortfall from the remaining trials
        S = pd.concat([S, L.drop(S.index).sample(n=n_target - len(S), random_state=SEED)])
    S = S.sort_values("nct_id").reset_index(drop=True)
    S["us_site"] = "Y"
    S["in_scope_condition"] = "Y"
    S["sponsor_class"] = S.sponsor_class.where(S.sponsor_class == "INDUSTRY", "other")
    S["stop_category"] = S.stop_category.where(S.overall_status.isin(STOPPED), "n/a")
    S["url"] = "https://clinicaltrials.gov/study/" + S.nct_id
    S[["nct_id", "stratum", "url"] + [c for c, _, _, _ in CHECKS]].to_csv(P / "handcheck_sample.csv", index=False)

    wb = Workbook()
    # --- Instructions ---------------------------------------------------------
    ws = wb.active
    ws.title = "Instructions"
    lines = [
        ("Hand check: Cardiometabolic Trial Accrual Ledger v1.0", True),
        (f"{len(S)} trials, drawn at random (10 percent of {len(L)}; seed {SEED}; "
         "proportional to completed / stopped / ongoing).", False),
        ("", False),
        ("How to use this form", True),
        ("1. Go to the 'Check' tab. Each row is one trial. Click the link in column B to open its registry page.", False),
        ("2. Grey cells show what the ledger says. Do not change them.", False),
        ("3. In the YELLOW cell to the right of each grey cell, pick what the registry page says from the", False),
        ("   dropdown list. The list uses the ledger's codes; the 'Codes' tab shows which code matches", False),
        ("   each wording on the website (for example, 'Active, not recruiting' is ACTIVE_NOT_RECRUITING).", False),
        ("   Enrollment is the only cell you type: enter the number shown, without commas.", False),
        ("4. The 'Match' cell then shows OK or CHECK. For every CHECK, add a short note in the last column.", False),
        ("5. Row 2 of the Check tab is a worked example (fake trial). Leave it; it is not counted.", False),
        ("6. The 'Summary' tab updates by itself and shows the agreement rate for each field.", False),
        ("", False),
        ("Where to look on each registry page", True),
    ] + [(f"- {label}: {where}", False) for _, label, where, _ in CHECKS] + [
        ("", False),
        ("Planned enrollment is not checked here; it arrives in v1.1 with its own short sample.", False),
        ("Time: about 2 minutes per trial, roughly 9 to 10 hours in total. Doing 50 a sitting works well.", False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=F, bold=bold, size=12 if i == 1 else 11)
    ws.column_dimensions["A"].width = 120

    # --- Check sheet ------------------------------------------------------------
    cs = wb.create_sheet("Check")
    headers = ["NCT ID", "Registry link", "Group"]
    for _, label, _, _ in CHECKS:
        headers += [f"Ledger: {label}", "Registry says", "Match"]
    headers += ["Notes"]
    for j, h in enumerate(headers, 1):
        c = cs.cell(row=1, column=j, value=h)
        c.font = Font(name=F, bold=True)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    cs.row_dimensions[1].height = 60
    cs.freeze_panes = "B2"

    example = ["NCT00000000 (EXAMPLE)", "https://clinicaltrials.gov/", "completed"]
    ex_vals = ["COMPLETED", 120, "ACTUAL", "PHASE3", "INDUSTRY", "Y", "Y", "n/a"]
    ex_reg = ["COMPLETED", 118, "ACTUAL", "PHASE3", "INDUSTRY", "Y", "Y", "n/a"]

    def write_row(r, ident, url, grp, vals, reg, is_example=False):
        cs.cell(row=r, column=1, value=ident).font = Font(name=F, italic=is_example)
        lc = cs.cell(row=r, column=2, value=url)
        if not is_example:
            lc.hyperlink = url
            lc.font = Font(name=F, color="0563C1", underline="single")
        else:
            lc.font = Font(name=F, italic=True)
        cs.cell(row=r, column=3, value=grp).font = Font(name=F)
        for k, v in enumerate(vals):
            col = 4 + 3 * k
            a = cs.cell(row=r, column=col, value=v)
            a.fill, a.font = GREY, Font(name=F)
            b = cs.cell(row=r, column=col + 1, value=reg[k] if reg else None)
            b.fill, b.font = YELLOW, Font(name=F, italic=is_example)
            al, bl = get_column_letter(col), get_column_letter(col + 1)
            m = cs.cell(row=r, column=col + 2,
                        value=f'=IF({bl}{r}="","",IF(UPPER(TRIM({al}{r}&""))=UPPER(TRIM({bl}{r}&"")),"OK","CHECK"))')
            m.font = Font(name=F, bold=True)
        cs.cell(row=r, column=len(headers),
                value="Example: enrollment differs; registry shows 118" if is_example else None).font = Font(name=F)

    write_row(2, *example, ex_vals, ex_reg, is_example=True)
    for i, row in S.iterrows():
        write_row(3 + i, row.nct_id, row.url, row.stratum, [row[c] for c, _, _, _ in CHECKS], None)
    last = 2 + len(S)
    widths = [14, 44, 18] + [16, 14, 9] * len(CHECKS) + [40]
    for j, w in enumerate(widths, 1):
        cs.column_dimensions[get_column_letter(j)].width = w
    # --- Codes sheet: dropdown sources and the website wording for each code ----
    cd = wb.create_sheet("Codes")
    cd.cell(row=1, column=1, value="What each dropdown code means on the registry page").font = Font(name=F, bold=True, size=12)
    col = 1
    for c, label, _, codes in CHECKS:
        if not codes:
            continue
        h1 = cd.cell(row=2, column=col, value=f"{label}: code")
        h2 = cd.cell(row=2, column=col + 1, value="Website shows")
        for h in (h1, h2):
            h.font, h.fill = Font(name=F, bold=True), HEAD
        for i, (code, words) in enumerate(codes):
            cd.cell(row=3 + i, column=col, value=code).font = Font(name=F)
            cd.cell(row=3 + i, column=col + 1, value=words).font = Font(name=F)
        letter = get_column_letter(col)
        rng = f"Codes!${letter}$3:${letter}${2 + len(codes)}"
        k = [x[0] for x in CHECKS].index(c)
        target = get_column_letter(5 + 3 * k)
        dv = DataValidation(type="list", formula1=f"={rng}", allow_blank=True, showErrorMessage=True,
                            errorTitle="Pick from the list",
                            error="Choose a code from the dropdown. The Codes tab shows what each code means.")
        cs.add_data_validation(dv)
        dv.add(f"{target}2:{target}{last}")
        cd.column_dimensions[letter].width = 26
        cd.column_dimensions[get_column_letter(col + 1)].width = 40
        col += 3
    k = [x[0] for x in CHECKS].index("enrollment_reported")
    target = get_column_letter(5 + 3 * k)
    dvn = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                         showErrorMessage=True, errorTitle="Whole number",
                         error="Type the enrollment number only, with no commas or words.")
    cs.add_data_validation(dvn)
    dvn.add(f"{target}2:{target}{last}")

    # --- Summary ----------------------------------------------------------------
    sm = wb.create_sheet("Summary")
    for j, h in enumerate(["Field", "Checked", "OK", "CHECK", "Agreement"], 1):
        c = sm.cell(row=1, column=j, value=h)
        c.font, c.fill = Font(name=F, bold=True), HEAD
    for k, (_, label, _, _) in enumerate(CHECKS):
        r = 2 + k
        mc = get_column_letter(6 + 3 * k)
        rng = f"Check!{mc}3:{mc}{last}"
        sm.cell(row=r, column=1, value=label)
        sm.cell(row=r, column=2, value=f'=COUNTIF({rng},"OK")+COUNTIF({rng},"CHECK")')
        sm.cell(row=r, column=3, value=f'=COUNTIF({rng},"OK")')
        sm.cell(row=r, column=4, value=f'=COUNTIF({rng},"CHECK")')
        sm.cell(row=r, column=5, value=f'=IF(B{r}=0,"",C{r}/B{r})').number_format = "0.0%"
        for j in range(1, 6):
            sm.cell(row=r, column=j).font = Font(name=F)
    r = 3 + len(CHECKS)
    yellow_cols = [get_column_letter(5 + 3 * k) for k in range(len(CHECKS))]
    filled = "*".join(f'(Check!{c}3:{c}{last}<>"")' for c in yellow_cols)
    sm.cell(row=r, column=1, value="Trials with every yellow cell filled in").font = Font(name=F, bold=True)
    sm.cell(row=r, column=2, value=f"=SUMPRODUCT({filled})").font = Font(name=F, bold=True)
    sm.cell(row=r + 1, column=1, value=f"Trials in sample: {len(S)} (example row 2 excluded from all counts)").font = Font(name=F, italic=True)
    sm.column_dimensions["A"].width = 50
    for col in "BCDE":
        sm.column_dimensions[col].width = 12

    out = ROOT / "report" / "handcheck_form.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print(f"sample {len(S)} trials:", S.stratum.value_counts().to_dict(), "->", out)


if __name__ == "__main__":
    main()
